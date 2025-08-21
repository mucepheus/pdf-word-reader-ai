import os
import sys
import tempfile
from docx import Document

sys.path.append(os.path.dirname(os.path.dirname(__file__)))

from document_reader import DocumentReader
from qa_engine import QAEngine

def test_docx_question_answering():
    doc = Document()
    doc.add_paragraph("The patient has a fever and cough.")
    doc.add_paragraph("Treatment includes rest and fluids.")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)

    reader = DocumentReader()
    texts = reader.read(tmp.name)
    engine = QAEngine(texts)
    answer = engine.ask("What symptoms does the patient have?")

    os.unlink(tmp.name)
    assert "fever" in answer.lower()
    assert "cough" in answer.lower()
