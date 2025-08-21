import io
from typing import List

try:
    import pdfplumber
except ImportError:  # pragma: no cover - optional dependency
    pdfplumber = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover
    Image = None

try:
    import pytesseract
except ImportError:  # pragma: no cover
    pytesseract = None

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


class DocumentReader:
    """Read text, tables, and images from various document formats."""

    def read(self, path: str) -> List[str]:
        path_lower = path.lower()
        if path_lower.endswith(".pdf"):
            return self._read_pdf(path)
        if path_lower.endswith(".docx"):
            return self._read_docx(path)
        if path_lower.endswith(".txt"):
            return self._read_txt(path)
        if path_lower.endswith(".xlsx"):
            return self._read_xlsx(path)
        raise ValueError(f"Unsupported file format: {path}")

    # ------------------------------------------------------------------
    def _read_pdf(self, path: str) -> List[str]:
        if pdfplumber is None:
            raise ImportError("pdfplumber is required to read PDF files")
        texts: List[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text:
                    texts.append(text)
                # tables
                for table in page.extract_tables():
                    table_text = "\n".join(
                        ["\t".join(cell or "" for cell in row) for row in table]
                    )
                    texts.append(table_text)
                # images with OCR
                if pytesseract and Image:
                    for img in page.images:
                        try:
                            bbox = (img["x0"], img["top"], img["x1"], img["bottom"])
                            cropped = page.crop(bbox).to_image(resolution=300)
                            ocr_text = pytesseract.image_to_string(cropped.original)
                            if ocr_text.strip():
                                texts.append(ocr_text)
                        except Exception:
                            continue
        return texts

    # ------------------------------------------------------------------
    def _read_docx(self, path: str) -> List[str]:
        if Document is None:
            raise ImportError("python-docx is required to read DOCX files")
        doc = Document(path)
        texts: List[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                texts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                texts.append("\t".join(cell.text for cell in row.cells))
        if pytesseract and Image:
            for rel in doc.part.rels.values():
                if "image" in rel.reltype:
                    try:
                        data = rel.target_part.blob
                        img = Image.open(io.BytesIO(data))
                        ocr_text = pytesseract.image_to_string(img)
                        if ocr_text.strip():
                            texts.append(ocr_text)
                    except Exception:
                        continue
        return texts

    # ------------------------------------------------------------------
    def _read_txt(self, path: str) -> List[str]:
        with open(path, "r", encoding="utf-8") as f:
            return [line.strip() for line in f if line.strip()]

    # ------------------------------------------------------------------
    def _read_xlsx(self, path: str) -> List[str]:
        if openpyxl is None:
            raise ImportError("openpyxl is required to read XLSX files")
        wb = openpyxl.load_workbook(path)
        texts: List[str] = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                if any(cell is not None and str(cell).strip() for cell in row):
                    texts.append("\t".join("" if cell is None else str(cell) for cell in row))
        return texts
